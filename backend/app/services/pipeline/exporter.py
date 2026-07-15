"""Document Export Service - Generates filled PDF/DOCX output.

Takes the pipeline results (detected fields + filled values)
and produces a professional output document.

Supports:
- PDF with filled form fields (interactive forms)
- PDF with text overlay (for scanned docs)
- DOCX with replaced placeholders
"""

import logging
import os
import uuid
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF

from app.core.config import settings

logger = logging.getLogger(__name__)


class DocumentExportService:
    """Exports filled documents as PDF or DOCX."""

    def __init__(self):
        self.export_dir = Path(settings.UPLOAD_DIR) / "exports"
        self.export_dir.mkdir(parents=True, exist_ok=True)

    async def export_pdf(
        self,
        source_path: str,
        filled_fields: list[dict[str, Any]],
    ) -> dict[str, Any]:
        """Export a filled PDF document.

        Strategy:
        1. If source is PDF with form widgets → fill widgets directly
        2. If source is PDF without widgets → add text overlays
        3. If source is image → create new PDF with text
        """
        ext = Path(source_path).suffix.lower()

        if ext == ".pdf":
            return await self._fill_pdf(source_path, filled_fields)
        elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".webp"):
            return await self._image_to_filled_pdf(source_path, filled_fields)
        else:
            # For DOCX or other, convert to PDF approach
            return await self._create_filled_pdf(filled_fields)

    async def _fill_pdf(
        self, source_path: str, filled_fields: list[dict[str, Any]]
    ) -> dict[str, Any]:
        """Fill an existing PDF's form widgets or add text overlays."""
        doc = fitz.open(source_path)
        output_name = f"filled_{uuid.uuid4().hex[:8]}.pdf"
        output_path = str(self.export_dir / output_name)

        # Build lookup of field values
        values = {}
        for f in filled_fields:
            name = f.get("field_name") or f.get("label") or f.get("name", "")
            value = f.get("value") or f.get("suggested_value")
            if name and value:
                values[name.lower()] = str(value)

        # Try to fill interactive form widgets
        widgets_filled = 0
        for page in doc:
            for widget in page.widgets():
                widget_name = (widget.field_name or "").lower()
                widget_label = (widget.field_label or "").lower()
                # Match by name or label
                for key, value in values.items():
                    if key in widget_name or key in widget_label or widget_name in key:
                        widget.field_value = value
                        widget.update()
                        widgets_filled += 1
                        break

        # If no widgets were filled, add text as overlay on first page
        if widgets_filled == 0 and values:
            page = doc[0]
            y_pos = 50
            for name, value in list(values.items())[:20]:
                text = f"{name.title()}: {value}"
                page.insert_text(
                    fitz.Point(50, y_pos),
                    text,
                    fontsize=11,
                    fontname="helv",
                    color=(0, 0, 0.6),
                )
                y_pos += 18

        doc.save(output_path)
        doc.close()

        file_size = os.path.getsize(output_path)
        logger.info(f"PDF export: {output_name} ({widgets_filled} widgets filled, {file_size} bytes)")

        return {
            "filename": output_name,
            "path": output_path,
            "size": file_size,
            "format": "pdf",
            "fields_filled": widgets_filled or len(values),
        }

    async def _image_to_filled_pdf(
        self, image_path: str, filled_fields: list[dict[str, Any]]
    ) -> dict[str, Any]:
        """Create a PDF from an image with filled field values on a second page."""
        output_name = f"filled_{uuid.uuid4().hex[:8]}.pdf"
        output_path = str(self.export_dir / output_name)

        doc = fitz.open()

        # Page 1: The original image
        img_doc = fitz.open(image_path)
        if len(img_doc) > 0:
            # If it's an image, insert it as a page
            img_rect = fitz.Rect(0, 0, 595, 842)  # A4 size
            page = doc.new_page(width=595, height=842)
            page.insert_image(img_rect, filename=image_path)
        img_doc.close()

        # Page 2: Filled field values (clean summary)
        page = doc.new_page(width=595, height=842)
        y_pos = 50
        page.insert_text(fitz.Point(50, y_pos), "Extracted & Filled Data", fontsize=16, fontname="helv", color=(0, 0, 0.6))
        y_pos += 30
        page.insert_text(fitz.Point(50, y_pos), "=" * 60, fontsize=8, color=(0.5, 0.5, 0.5))
        y_pos += 20

        for f in filled_fields:
            name = f.get("field_name") or f.get("label") or f.get("name", "Field")
            value = f.get("value") or f.get("suggested_value") or "—"
            text = f"{name}: {value}"
            page.insert_text(fitz.Point(50, y_pos), text, fontsize=11, fontname="helv")
            y_pos += 18
            if y_pos > 780:
                page = doc.new_page(width=595, height=842)
                y_pos = 50

        doc.save(output_path)
        doc.close()

        return {
            "filename": output_name,
            "path": output_path,
            "size": os.path.getsize(output_path),
            "format": "pdf",
            "fields_filled": len(filled_fields),
        }

    async def _create_filled_pdf(
        self, filled_fields: list[dict[str, Any]]
    ) -> dict[str, Any]:
        """Create a new PDF with all filled field values."""
        output_name = f"filled_{uuid.uuid4().hex[:8]}.pdf"
        output_path = str(self.export_dir / output_name)

        doc = fitz.open()
        page = doc.new_page(width=595, height=842)

        y_pos = 50
        page.insert_text(fitz.Point(50, y_pos), "Document - Filled Fields", fontsize=16, fontname="helv", color=(0, 0, 0.6))
        y_pos += 35

        for f in filled_fields:
            name = f.get("field_name") or f.get("label") or f.get("name", "Field")
            value = f.get("value") or f.get("suggested_value") or ""
            if value:
                page.insert_text(fitz.Point(50, y_pos), f"{name}:", fontsize=10, fontname="helv", color=(0.3, 0.3, 0.3))
                page.insert_text(fitz.Point(200, y_pos), str(value), fontsize=11, fontname="helv", color=(0, 0, 0))
                y_pos += 20
                if y_pos > 780:
                    page = doc.new_page(width=595, height=842)
                    y_pos = 50

        doc.save(output_path)
        doc.close()

        return {
            "filename": output_name,
            "path": output_path,
            "size": os.path.getsize(output_path),
            "format": "pdf",
            "fields_filled": len([f for f in filled_fields if f.get("value")]),
        }

    async def export_docx(
        self, source_path: str, filled_fields: list[dict[str, Any]]
    ) -> dict[str, Any]:
        """Export as DOCX with filled values."""
        import re
        from docx import Document as DocxDocument

        output_name = f"filled_{uuid.uuid4().hex[:8]}.docx"
        output_path = str(self.export_dir / output_name)

        # Build values lookup
        values = {}
        for f in filled_fields:
            name = f.get("field_name") or f.get("label") or f.get("name", "")
            value = f.get("value") or f.get("suggested_value")
            if name and value:
                values[name.lower()] = str(value)

        ext = Path(source_path).suffix.lower()
        if ext in (".docx", ".doc"):
            # Fill existing DOCX
            doc = DocxDocument(source_path)
            for para in doc.paragraphs:
                for field_name, value in values.items():
                    # Replace underlines or placeholders
                    if field_name in para.text.lower():
                        for run in para.runs:
                            if "___" in run.text or "____" in run.text:
                                run.text = re.sub(r"_{3,}", value, run.text)
            doc.save(output_path)
        else:
            # Create new DOCX with field values
            doc = DocxDocument()
            doc.add_heading("Document - Filled Fields", level=1)
            for name, value in values.items():
                doc.add_paragraph(f"{name.title()}: {value}")
            doc.save(output_path)

        return {
            "filename": output_name,
            "path": output_path,
            "size": os.path.getsize(output_path),
            "format": "docx",
            "fields_filled": len(values),
        }


# Singleton
export_service = DocumentExportService()
